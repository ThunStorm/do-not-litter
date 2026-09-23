from __future__ import annotations

import platform
from pathlib import Path


def read_document(path: Path, *, asr_provider: str | None = None) -> tuple[str, list[dict]]:
    suffix = path.suffix.lower()
    if suffix == ".docx":
        from docx import Document

        document = Document(path)
        segments = [
            {"text": paragraph.text, "locator": {"paragraph": index + 1}}
            for index, paragraph in enumerate(document.paragraphs)
            if paragraph.text.strip()
        ]
        return "\n".join(segment["text"] for segment in segments), segments
    if suffix == ".pdf":
        from pypdf import PdfReader

        reader = PdfReader(path)
        segments = []
        for page_index, page in enumerate(reader.pages):
            text = page.extract_text() or ""
            if text.strip():
                segments.append({"text": text, "locator": {"page": page_index + 1}})
        return "\n".join(segment["text"] for segment in segments), segments
    if suffix in {".xlsx", ".xlsm"}:
        from openpyxl import load_workbook

        workbook = load_workbook(path, read_only=True, data_only=True)
        segments = []
        for sheet in workbook.worksheets:
            for row_index, row in enumerate(sheet.iter_rows(values_only=True), start=1):
                values = [str(value).strip() for value in row if value is not None]
                if values:
                    segments.append(
                        {
                            "text": "\t".join(values),
                            "locator": {"sheet": sheet.title, "row": row_index},
                        }
                    )
        return "\n".join(segment["text"] for segment in segments), segments
    if suffix in {".txt", ".md", ".html", ".htm"}:
        text = path.read_text(encoding="utf-8", errors="replace")
        return text, [{"text": text, "locator": {"file": path.name}}]
    if suffix in {".png", ".jpg", ".jpeg"}:
        from zhijian.providers.ocr import MacVisionOCRProvider, TesseractOCRProvider

        provider = MacVisionOCRProvider() if platform.system() == "Darwin" else TesseractOCRProvider()
        try:
            return provider.recognize(path)
        except Exception as exc:
            return "", [
                {
                    "text": "",
                    "locator": {"file": path.name, "ocr_required": True, "error": str(exc)[:500]},
                    "confidence": None,
                }
            ]
    if suffix in {".mp3", ".m4a", ".wav", ".aac", ".mp4", ".mov", ".webm"}:
        from zhijian.ai.resource_manager import local_ai_resource_manager
        from zhijian.core.config import get_settings
        from zhijian.providers.asr import ASRProviderRegistry

        settings = get_settings()
        registry = ASRProviderRegistry(
            settings.whisper_binary,
            settings.whisper_model,
            settings.qwen_asr_python if settings.qwen_asr_enabled else None,
            settings.qwen_asr_runner if settings.qwen_asr_enabled else None,
            settings.qwen_asr_model if settings.qwen_asr_enabled else None,
            settings.qwen_asr_aligner_model if settings.qwen_asr_enabled else None,
            settings.qwen_asr_timeout_seconds,
        )
        provider = registry.get(asr_provider or settings.default_asr_provider)
        return local_ai_resource_manager.run("ASR", lambda: provider.transcribe(path))
    raise ValueError(f"暂不支持的文件类型：{suffix or 'unknown'}")
